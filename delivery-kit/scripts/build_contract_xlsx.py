# -*- coding: utf-8 -*-
"""生成 02-签约功能清单.xlsx。

用法（在 backend/ 下执行，经 uv 临时环境安装依赖，不污染 pyproject.toml）：
    cd backend
    uv run --with openpyxl python ..\\delivery-kit\\scripts\\build_contract_xlsx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from data_contract import (  # noqa: E402
    APIS,
    DEMO_SCRIPT,
    DEPLOY,
    DEPLOY_LIMITS,
    EXCLUSIONS,
    FUNCTIONS,
    GEN_DATE,
    LIMITS,
    VERSION,
)

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "delivery-kit" / "02-签约功能清单.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
NOTE_FONT = Font(size=9, color="595959")
ALT_FILL = PatternFill("solid", fgColor="F2F5FA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, headers, rows, widths, start_row=1, freeze=True):
    """写入表头与数据，套隔行底色 + 边框 + 自动换行。"""
    for idx, head in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=idx, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    for r_off, row in enumerate(rows, start=start_row + 1):
        for c_off, val in enumerate(row, start=1):
            cell = ws.cell(row=r_off, column=c_off, value=val)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
            if (r_off - start_row) % 2 == 0:
                cell.fill = ALT_FILL
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[start_row].height = 30
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    )


def build_cover(ws):
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "GraphRAG-Agent 多模态企业知识库平台"
    ws["A1"].font = Font(bold=True, size=18, color="1F3864")
    ws.merge_cells("A1:B1")
    ws["A2"] = "签约功能清单（合同附件 / 验收依据）"
    ws["A2"].font = Font(bold=True, size=13)
    ws.merge_cells("A2:B2")

    meta = [
        ("文档版本", VERSION),
        ("生成日期", GEN_DATE),
        ("事实来源", "本仓库代码；未经代码核实的规划口径一律不收录"),
        ("关联文档", "01-宣传素材包.md（对外宣传口径）；03-技术规格说明书.md（技术评审口径）"),
    ]
    row = 4
    for key, val in meta:
        ws.cell(row=row, column=1, value=key).font = Font(bold=True, size=10)
        c = ws.cell(row=row, column=2, value=val)
        c.alignment = WRAP_TOP
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="统计").font = TITLE_FONT
    row += 1
    stats = [
        ("交付功能项", f"{len(FUNCTIONS)} 项（覆盖 12 个功能模块）"),
        ("其中可现场演示", f"{sum(1 for f in FUNCTIONS if f[5].startswith('可演示'))} 项"),
        ("对外接口", f"{len(APIS)} 个 REST 端点"),
        ("验收演示步骤", f"{len(DEMO_SCRIPT)} 步"),
        ("明确不含项", f"{len(EXCLUSIONS)} 项（**合同条款须引用本清单**）"),
        ("已知限制", f"{len(LIMITS)} 条"),
    ]
    for key, val in stats:
        ws.cell(row=row, column=1, value=key).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2, value=val).alignment = WRAP_TOP
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="口径声明").font = TITLE_FONT
    row += 1
    notes = [
        "1. 本清单每一条均可回溯到代码落点（见各表最后一列的落点与限定语），不含营销修辞。",
        "2. 「可演示」指可在标准演示环境（已关闭 Mock、外部服务可用）中现场操作并复现。",
        "3. 「限定语」是功能边界的组成部分，对外讲解时须与功能描述同时告知，不得单独隐去。",
        "4. 本清单**不含**的部分请见「03-明确不含」工作表，建议作为合同除外条款直接引用。",
        "5. 演示语料为虚构数据，任何识别结果均不构成对真实业务的结论。",
        "6. 评测类数据基于受控演示语料且阈值尚未校准，不得表述为正式达标结论。",
    ]
    for note in notes:
        c = ws.cell(row=row, column=2, value=note)
        c.font = NOTE_FONT
        c.alignment = WRAP_TOP
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="验收签署").font = TITLE_FONT
    row += 1
    for role in ("甲方代表", "乙方代表", "签署日期"):
        ws.cell(row=row, column=1, value=role).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2, value="________________________").alignment = WRAP_TOP
        row += 1


def build_functions(ws):
    headers = [
        "编号", "功能模块", "功能项", "功能说明（可验收事实）",
        "交付形态", "可否演示", "验收方式", "验收判据（通过标准）",
        "前端落点", "后端落点", "限定语 / 边界（须同时告知）",
    ]
    widths = [8, 16, 20, 46, 14, 10, 34, 48, 30, 34, 44]
    write_table(ws, headers, FUNCTIONS, widths)


def build_demo(ws):
    headers = ["步骤", "序号", "操作动作", "操作页面", "预期看到的结果",
               "对应功能编号", "需准备数据", "通过标准"]
    widths = [8, 8, 32, 20, 46, 16, 12, 40]
    write_table(ws, headers, DEMO_SCRIPT, widths)


def build_exclusions(ws):
    headers = ["编号", "明确不含的能力", "为什么容易被误读（原因）", "真实情况说明"]
    widths = [10, 40, 40, 76]
    write_table(ws, headers, EXCLUSIONS, widths)


def build_apis(ws):
    headers = ["方法", "路径（前缀 /api/v1）", "功能说明", "前端是否已接入", "是否在演示流程中"]
    widths = [10, 44, 44, 16, 18]
    write_table(ws, headers, APIS, widths)


def build_deploy(ws):
    headers = ["类别", "组件", "技术选型", "交付镜像 / 形态", "端口"]
    widths = [14, 22, 52, 40, 16]
    write_table(ws, headers, DEPLOY, widths)

    start = len(DEPLOY) + 4
    ws.cell(row=start, column=1, value="部署约束与注意事项").font = TITLE_FONT
    write_table(
        ws,
        ["类别", "约束说明"],
        DEPLOY_LIMITS,
        [14, 22],
        start_row=start + 1,
        freeze=False,
    )
    ws.column_dimensions["B"].width = 110


def build_limits(ws):
    headers = ["编号", "类别", "限制项", "影响说明", "应对建议"]
    widths = [10, 14, 40, 48, 44]
    write_table(ws, headers, LIMITS, widths)


def main() -> int:
    wb = Workbook()

    ws = wb.active
    ws.title = "00-封面与口径"
    build_cover(ws)

    builders = [
        ("01-功能清单", build_functions),
        ("02-演示验收剧本", build_demo),
        ("03-明确不含", build_exclusions),
        ("04-API接口清单", build_apis),
        ("05-部署与运行环境", build_deploy),
        ("06-已知限制", build_limits),
    ]
    for name, fn in builders:
        fn(wb.create_sheet(name))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"[OK] 已生成: {OUT}")
    print(f"     功能项 {len(FUNCTIONS)} | 接口 {len(APIS)} | "
          f"剧本 {len(DEMO_SCRIPT)} 步 | 不含项 {len(EXCLUSIONS)} | 限制 {len(LIMITS)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
