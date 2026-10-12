# -*- coding: utf-8 -*-
"""轻量 Markdown → DOCX 转换器（仅覆盖本项目交付文档用到的语法子集）。

支持的语法：
  - 标题 `#` ~ `####`
  - 表格（含分隔行），自动识别表头对齐分割线
  - 围栏代码块 ```
  - 引用块 `> `
  - 无序列表 `- ` 与有序列表 `1. `
  - 分隔线 `---`
  - 行内 `**粗体**` 与 `` `行内代码` ``
  - 普通段落（自动合并连续行）

字体：正文中文用「微软雅黑」，代码用「Consolas」，并设置东亚字体以避免中文回退。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)

RE_H = re.compile(r"^(#{1,4})\s+(.*)$")
RE_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
RE_CODE = re.compile(r"^```")
RE_QUOTE = re.compile(r"^>\s?(.*)$")
RE_UL = re.compile(r"^[-*]\s+(.*)$")
RE_OL = re.compile(r"^(\d+)\.\s+(.*)$")
RE_HR = re.compile(r"^(-{3,}|\*{3,})$")


def set_cn_font(run, name=CN_FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def add_inline(par, text, *, code=False, size=None):
    """把一段文本按 **粗体** 与 `代码` 拆分成多个 run。"""
    if code:
        run = par.add_run(text)
        set_cn_font(run, CODE_FONT)
        run.font.size = Pt(size or 9)
        return
    tokens = re.findall(r"(\*\*|`)(.+?)\1|([^`*]+)", text, flags=re.S)
    for kind, captured, plain in tokens:
        if kind:  # 命中了 ** 或 ` 包裹的片段
            piece = captured
            run = par.add_run(piece)
            set_cn_font(run, CODE_FONT if kind == "`" else CN_FONT)
            run.bold = kind == "**"
            if kind == "`":
                run.font.color.rgb = RGBColor(0xA3, 0x18, 0x18)
            run.font.size = Pt(size or 10.5)
        else:  # 普通文本
            piece = plain
            if not piece:
                continue
            run = par.add_run(piece)
            set_cn_font(run)
            run.font.size = Pt(size or 10.5)


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    # 忽略表格内的转义竖线
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def add_table(doc, rows: list[str]) -> None:
    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]
    # 补齐列数，防止 Markdown 表格书写不齐导致越界
    ncol = max([len(header)] + [len(r) for r in body])
    header += [""] * (ncol - len(header))
    body = [r + [""] * (ncol - len(r)) for r in body]

    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(text.replace("**", ""))
        set_cn_font(run)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1F3864")

    for ridx, row in enumerate(body):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, text, size=9)
            if ridx % 2 == 1:
                shade(cells[i], "F2F5FA")


def shade(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def convert(md_path: Path, docx_path: Path, *, title: str | None = None) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    if title:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(title)
        set_cn_font(run)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = ACCENT

    i = 0
    paragraph_buf: list[str] = []
    n = len(lines)

    def flush_paragraph():
        nonlocal paragraph_buf
        if not paragraph_buf:
            return
        text = " ".join(x.strip() for x in paragraph_buf).strip()
        paragraph_buf = []
        if not text:
            return
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.line_spacing = 1.4
        add_inline(par, text)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if RE_CODE.match(stripped):
            flush_paragraph()
            i += 1
            buf = []
            while i < n and not RE_CODE.match(lines[i].strip()):
                buf.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.5)
            par.paragraph_format.space_after = Pt(8)
            run = par.add_run("\n".join(buf))
            set_cn_font(run, CODE_FONT)
            run.font.size = Pt(8.5)
            continue

        if RE_H.match(stripped):
            flush_paragraph()
            level = len(RE_H.match(stripped).group(1))
            text = RE_H.match(stripped).group(2).replace("**", "")
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            par.paragraph_format.space_after = Pt(6)
            run = par.add_run(text)
            set_cn_font(run)
            run.bold = True
            run.font.size = Pt({1: 16, 2: 13.5, 3: 12, 4: 11}[level])
            if level <= 3:
                run.font.color.rgb = ACCENT
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            if len(block) >= 2 and RE_TABLE_SEP.match(block[1].strip()):
                add_table(doc, block)
            else:
                for b in block:
                    par = doc.add_paragraph()
                    add_inline(par, b, size=9)
            continue

        if RE_HR.match(stripped):
            flush_paragraph()
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(4)
            ppr = par._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BFBFBF")
            pbdr.append(bottom)
            ppr.append(pbdr)
            i += 1
            continue

        if RE_QUOTE.match(stripped):
            flush_paragraph()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.paragraph_format.space_after = Pt(6)
            run = par.add_run(RE_QUOTE.match(stripped).group(1))
            set_cn_font(run)
            run.font.size = Pt(9.5)
            run.italic = True
            run.font.color.rgb = GREY
            i += 1
            continue

        if RE_UL.match(stripped):
            flush_paragraph()
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, RE_UL.match(stripped).group(1))
            i += 1
            continue

        if RE_OL.match(stripped):
            flush_paragraph()
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, RE_OL.match(stripped).group(2))
            i += 1
            continue

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        paragraph_buf.append(stripped)
        i += 1

    flush_paragraph()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"[OK] 已生成: {docx_path}")


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        raise SystemExit("用法: python md_to_docx.py <input.md> <output.docx> [title]")
    convert(
        Path(sys.argv[1]),
        Path(sys.argv[2]),
        title=sys.argv[3] if len(sys.argv) > 3 else None,
    )
